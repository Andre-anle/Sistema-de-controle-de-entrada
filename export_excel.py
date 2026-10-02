from datetime import datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

VERDE = "1F6B3A"
VERDE_CLARO = "2F8A4A"
VERDE_SUAVE = "E7F3EA"
CREME = "F3F0E6"
LARANJA = "E07A2F"
BRANCO = "FFFFFF"
TEXTO = "243028"
TEXTO_SUAVE = "5B6F60"
BORDA = "C9D6C3"

COLUNAS_MERCADORIAS = (
    ("EAN", "ean", "texto"),
    ("Mercadoria", "descricao", "texto"),
    ("Colaborador", "colaborador", "texto"),
    ("Qtd", "quantidade", "numero"),
    ("Data e hora", "data_hora", "data"),
)
COLUNAS_VISITANTES = (
    ("Visitante", "nome", "texto"),
    ("Documento", "documento", "texto"),
    ("Função", "funcao", "texto"),
    ("Empresa", "empresa", "texto"),
    ("Autorizado por", "autorizado_por", "texto"),
    ("Entrada", "entrada", "data"),
    ("Saída", "saida", "data"),
)

_FILL_CREME = PatternFill("solid", fgColor=CREME)
_FILL_VERDE = PatternFill("solid", fgColor=VERDE)
_FILL_VERDE_CLARO = PatternFill("solid", fgColor=VERDE_CLARO)
_FILL_ZEBRA = PatternFill("solid", fgColor=VERDE_SUAVE)

_FONTE_MARCA = Font(name="Calibri", bold=True, size=16, color=VERDE)
_FONTE_TITULO = Font(name="Calibri", bold=True, size=13, color=LARANJA)
_FONTE_FILTRO = Font(name="Calibri", italic=True, size=10, color=TEXTO_SUAVE)
_FONTE_CABECALHO = Font(name="Calibri", bold=True, size=11, color=BRANCO)
_FONTE_VAZIO = Font(name="Calibri", italic=True, size=11, color=TEXTO_SUAVE)
_FONTE_TOTAL = Font(name="Calibri", bold=True, size=11, color=BRANCO)

_ALINHAMENTO = Alignment(vertical="center")
_ALINHAMENTO_QUEBRA = Alignment(vertical="center", wrap_text=True)
_ALINHAMENTO_CENTRO = Alignment(horizontal="center", vertical="center")
_ALINHAMENTO_CABECALHO = Alignment(horizontal="center", vertical="center")

_BORDA = Border(
    left=Side(style="thin", color=BORDA),
    right=Side(style="thin", color=BORDA),
    top=Side(style="thin", color=BORDA),
    bottom=Side(style="thin", color=BORDA),
)

_FORMATO_DATA = "dd/mm/yyyy hh:mm:ss"
_ESTILO_TABELA = TableStyleInfo(
    name="TableStyleMedium7",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)


# Injeção de fórmula: um campo digitado como  =HYPERLINK(...)  ou  =cmd|'/c calc'!A1
# viraria fórmula ao abrir o Excel. Aqui o valor é sempre gravado como TEXTO.
_PREFIXOS_FORMULA = ("=", "+", "-", "@", "\t", "\r", "\n")


def _texto_seguro(texto: str) -> str:
    """Remove caracteres de controle que o formato xlsx não aceita (evita falha na exportação)."""
    return ILLEGAL_CHARACTERS_RE.sub("", texto)


def _neutralizar_formula(celula, texto: str) -> None:
    if texto.startswith(_PREFIXOS_FORMULA):
        celula.data_type = "s"  # nunca "f" (fórmula)
        celula.quotePrefix = True  # o Excel também mostra como texto ao editar


def _para_data(valor: Any) -> datetime | str:
    texto = str(valor or "").strip()
    if not texto:
        return ""
    try:
        return datetime.fromisoformat(texto)
    except ValueError:
        pass
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return texto


def _estilizar(celula, *, font=None, fill=None, alignment=None, border=None, number_format=None) -> None:
    if font is not None:
        celula.font = font
    if fill is not None:
        celula.fill = fill
    if alignment is not None:
        celula.alignment = alignment
    if border is not None:
        celula.border = border
    if number_format is not None:
        celula.number_format = number_format


def _largura_colunas(
    colunas: tuple[tuple[str, str, str], ...],
    linhas: list[Any],
) -> list[float]:
    larguras = [min(max(len(rotulo) + 4, 14), 48) for rotulo, _, _ in colunas]
    for registro in linhas:
        for indice, (_, chave, tipo) in enumerate(colunas):
            bruto = registro[chave]
            if tipo == "data":
                tamanho = 19
            else:
                tamanho = len(str(bruto)) if bruto is not None else 0
            candidato = min(max(tamanho + 4, 14), 48)
            if candidato > larguras[indice]:
                larguras[indice] = candidato
    return larguras


def _limitar_area(ws: Worksheet, qtd: int, ultima_linha: int, larguras: list[float]) -> None:
    letra = get_column_letter(qtd)
    ws.print_area = f"A1:{letra}{ultima_linha}"
    for indice, largura in enumerate(larguras, start=1):
        dimensao = ws.column_dimensions[get_column_letter(indice)]
        dimensao.width = largura



def _configurar_pagina(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.oddHeader.left.text = "Hortifruti Natural da Terra"
    ws.oddFooter.right.text = "Página &P de &N"
    ws.print_title_rows = "1:5"
    ws.sheet_format.defaultRowHeight = 18
    ws.freeze_panes = "A6"


def _escrever_cabecalho(ws: Worksheet, titulo: str, filtros: str, qtd: int, total: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=qtd)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=qtd)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=qtd)

    marca = ws.cell(1, 1, "HORTIFRUTI NATURAL DA TERRA")
    _estilizar(marca, font=_FONTE_MARCA, fill=_FILL_CREME, alignment=_ALINHAMENTO)

    titulo_celula = ws.cell(2, 1, titulo)
    _estilizar(titulo_celula, font=_FONTE_TITULO, fill=_FILL_CREME, alignment=_ALINHAMENTO)

    gerado = datetime.now().strftime("%d/%m/%Y %H:%M")
    filtro = ws.cell(
        3, 1, _texto_seguro(f"Gerado em {gerado}   ·   Filtros: {filtros}   ·   {total} registro(s)")
    )
    _estilizar(filtro, font=_FONTE_FILTRO, fill=_FILL_CREME, alignment=_ALINHAMENTO_QUEBRA)

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 32
    ws.row_dimensions[4].height = 8
    ws.row_dimensions[5].height = 22


def _escrever_aba(
    ws: Worksheet,
    titulo: str,
    filtros: str,
    colunas: tuple[tuple[str, str, str], ...],
    registros: Iterable[Any],
    nome_tabela: str,
) -> None:
    linhas = list(registros)
    qtd = len(colunas)
    letra = get_column_letter(qtd)
    _configurar_pagina(ws)
    _escrever_cabecalho(ws, titulo, filtros, qtd, len(linhas))
    larguras = _largura_colunas(colunas, linhas)

    for indice, (rotulo, _, _) in enumerate(colunas, start=1):
        celula = ws.cell(5, indice, rotulo)
        _estilizar(
            celula,
            font=_FONTE_CABECALHO,
            fill=_FILL_VERDE,
            alignment=_ALINHAMENTO_CABECALHO,
            border=_BORDA,
        )

    if not linhas:
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=qtd)
        vazio = ws.cell(6, 1, "Nenhum registro encontrado para os filtros selecionados.")
        _estilizar(
            vazio,
            font=_FONTE_VAZIO,
            fill=_FILL_ZEBRA,
            alignment=_ALINHAMENTO_CENTRO,
            border=_BORDA,
        )
        ws.row_dimensions[6].height = 24
        _limitar_area(ws, qtd, 6, larguras)
        return

    for offset, registro in enumerate(linhas):
        linha = 6 + offset
        for coluna, (_, chave, tipo) in enumerate(colunas, start=1):
            bruto = registro[chave]
            if tipo == "data":
                valor = _para_data(bruto)
            elif tipo == "numero":
                valor = bruto
            elif bruto is None:
                valor = None
            else:
                valor = _texto_seguro(str(bruto))
            celula = ws.cell(linha, coluna)
            celula.value = valor
            if isinstance(valor, str):
                _neutralizar_formula(celula, valor)
            if isinstance(valor, datetime):
                celula.number_format = _FORMATO_DATA

    ultima = 5 + len(linhas)
    tabela = Table(displayName=nome_tabela, ref=f"A5:{letra}{ultima}")
    tabela.tableStyleInfo = _ESTILO_TABELA
    ws.add_table(tabela)

    total_linha = ultima + 1
    ws.merge_cells(start_row=total_linha, start_column=1, end_row=total_linha, end_column=qtd)
    total = ws.cell(total_linha, 1, f"Total: {len(linhas)} registro(s)")
    _estilizar(total, font=_FONTE_TOTAL, fill=_FILL_VERDE_CLARO, alignment=_ALINHAMENTO, border=_BORDA)
    _limitar_area(ws, qtd, total_linha, larguras)


def gerar_planilha_entradas(
    mercadorias: list | None,
    visitantes: list | None,
    filtros_mercadorias: str,
    filtros_visitantes: str,
) -> bytes:
    wb = Workbook()
    padrao = wb.active
    primeira = True

    if mercadorias is not None:
        ws = padrao if primeira else wb.create_sheet()
        primeira = False
        ws.title = "Mercadorias"
        _escrever_aba(
            ws,
            "Entradas de mercadorias",
            filtros_mercadorias,
            COLUNAS_MERCADORIAS,
            mercadorias,
            "TabelaMercadorias",
        )
    if visitantes is not None:
        ws = padrao if primeira else wb.create_sheet()
        primeira = False
        ws.title = "Visitantes"
        _escrever_aba(
            ws,
            "Entradas de visitantes",
            filtros_visitantes,
            COLUNAS_VISITANTES,
            visitantes,
            "TabelaVisitantes",
        )
    if primeira:
        raise ValueError("Nenhum conteúdo selecionado para exportar.")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
